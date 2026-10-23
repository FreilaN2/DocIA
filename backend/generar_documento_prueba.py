# -*- coding: utf-8 -*-
"""
Script para generar el documento de prueba integral de Normas APA.
Incluye todos los elementos requeridos para validar:
1. Portada académica.
2. Jerarquía completa de títulos APA 7ma (Nivel 1 al Nivel 5).
3. Párrafos normales justificados con sangría de primera línea.
4. Citas cortas dentro de párrafos normales.
5. Citas en bloque largas (>= 40 palabras) con sangría de bloque.
6. Auditoría de Citas vs Bibliografía:
   - Citas válidas sincronizadas.
   - 2 Citas huérfanas intencionales: Martínez (2023) y López & Gómez (2021).
   - 2 Referencias fantasma intencionales: Zúñiga & Morales (2025) y Vargas (2020).
7. Lista de Referencias bibliográficas desordenada alfabéticamente para comprobar ordenación A-Z y sangría francesa.
"""

import os
import shutil
from docx import Document

def generar_documento_prueba():
    doc = Document()
    
    # ─── 1. PORTADA ACADÉMICA ──────────────────────────────────────
    doc.add_paragraph("UNIVERSIDAD CENTRAL DE VENEZUELA")
    doc.add_paragraph("FACULTAD DE CIENCIAS")
    doc.add_paragraph("ESCUELA DE COMPUTACIÓN")
    doc.add_paragraph("POSTGRADO EN CIENCIAS DE LA COMPUTACIÓN")
    doc.add_paragraph("")
    doc.add_paragraph("")
    doc.add_paragraph("APLICACIÓN DE MODELOS DE INTELIGENCIA ARTIFICIAL EN LA AUTOMATIZACIÓN DOCUMENTAL BAJO NORMAS APA")
    doc.add_paragraph("Trabajo Especial de Grado presentado como requisito parcial para optar al título de Magíster en Ciencias de la Computación")
    doc.add_paragraph("")
    doc.add_paragraph("")
    doc.add_paragraph("Autor: Ing. Carlos Eduardo Mendoza")
    doc.add_paragraph("Tutor Académico: Dr. Alejandro Villarroel")
    doc.add_paragraph("Caracas, Noviembre de 2026")
    
    # ─── 2. RESUMEN Y ABSTRACT (NIVEL 1) ──────────────────────────
    doc.add_paragraph("Resumen")
    doc.add_paragraph(
        "La presente investigación analiza la transformación operativa derivada de la adopción de "
        "arquitecturas basadas en modelos de lenguaje de gran escala (LLM) dentro de los flujos de trabajo "
        "académicos y editoriales durante el período 2025-2026. El estudio se enmarca en una metodología "
        "cuantitativa con diseño cuasi-experimental, evaluando el formateo automatizado de documentos "
        "complejos según los estándares oficiales de las Normas APA séptima edición. Los resultados demostraron "
        "una reducción del 84.5% en los tiempos de corrección tipográfica, una eliminación del 96% de errores "
        "en sangrías y alineaciones, y una sincronización perfecta entre las citas textuales y el listado de "
        "referencias bibliográficas. Se concluye que la inteligencia artificial generativa complementa de manera "
        "eficiente la labor docente e investigativa en la educación superior moderna. "
        "Palabras clave: inteligencia artificial, normas APA, automatización documental, procesamiento de lenguaje natural."
    )
    
    doc.add_paragraph("Abstract")
    doc.add_paragraph(
        "This research examines the operational transformation resulting from the adoption of Large "
        "Language Model (LLM) architectures within academic and editorial workflows during the 2025-2026 period. "
        "The study uses a quantitative quasi-experimental design to evaluate the automated formatting of complex "
        "manuscripts according to official APA 7th edition standards. The findings reveal an 84.5% reduction "
        "in typographical review times and a 96% elimination of indentation and alignment errors. "
        "Keywords: artificial intelligence, APA standards, document automation, natural language processing."
    )
    
    # ─── 3. CAPÍTULO I: EL PROBLEMA (NIVEL 1 Y 2) ──────────────────
    doc.add_paragraph("Capítulo I: El Problema de Investigación")
    doc.add_paragraph("Planteamiento del Problema")
    doc.add_paragraph(
        "En el contexto de la educación superior contemporánea, la redacción científica exige una rigurosa "
        "adherencia a estándares internacionales de formato y estilo editorial. Diversos estudios destacan "
        "que los investigadores dedican una cantidad sustancial de tiempo a resolver inconsistencias de sangrías, "
        "márgenes e interlineados en lugar de concentrarse en la profundidad conceptual de sus hallazgos."
    )
    
    # Párrafo normal con CITA CORTA en el texto (< 40 palabras entre comillas)
    doc.add_paragraph(
        "Al respecto, Navarro (2021) argumenta que \"el formateo manual de citas y referencias bibliográficas "
        "representa más del 40% de las observaciones formales en las defensas de tesis\" (p. 33). Esta dificultad "
        "se agrava por la dispersión de criterios entre las distintas versiones de las normas y la falta de "
        "herramientas informáticas adaptadas al contexto hispanohablante."
    )
    
    # Párrafo previo a la cita larga
    doc.add_paragraph(
        "La literatura metodológica más reciente advierte sobre la fatiga cognitiva y el impacto negativo "
        "que sufren los estudiantes de posgrado cuando deben formatear manualmente monografías de gran volumen, "
        "tal como lo describen de forma contundente Hernández et al. (2022):"
    )
    
    # CITA LARGA EN BLOQUE (>= 40 palabras)
    doc.add_paragraph(
        "La ausencia de mecanismos computacionales para la estandarización documental genera un esfuerzo cognitivo "
        "desproporcionado en los investigadores noveles. Cuando los tesistas dedican semanas a corregir márgenes, "
        "sangrías y referencias de manera artesanal, se reduce significativamente la calidad analítica y metodológica "
        "del contenido científico final en las universidades contemporáneas. (Hernández et al., 2022, p. 184)"
    )
    
    doc.add_paragraph(
        "Lo expuesto anteriormente justifica la necesidad de crear asistentes inteligentes capaces de procesar "
        "documentos estructurados en tiempo real garantizando el cumplimiento irrestricto de las normas académicas."
    )
    
    doc.add_paragraph("Objetivos de la Investigación")
    doc.add_paragraph("1.1 Objetivo General")
    doc.add_paragraph(
        "Desarrollar y evaluar un sistema de automatización documental basado en modelos de lenguaje y reglas "
        "heurísticas para la estandarización de manuscritos académicos bajo las Normas APA séptima edición."
    )
    doc.add_paragraph("1.2 Objetivos Específicos")
    doc.add_paragraph(
        "Diagnosticar los errores más frecuentes en la aplicación de sangrías, jerarquías de títulos y referencias "
        "en los trabajos de grado universitarios."
    )
    doc.add_paragraph(
        "Diseñar un motor híbrido que combine modelos de lenguaje avanzados con algoritmos deterministas para la "
        "detección de inconsistencias entre citas textuales y bibliografía."
    )
    
    # ─── 4. CAPÍTULO II: MARCO TEÓRICO (N1 A N5) ───────────────────
    doc.add_paragraph("Capítulo II: Marco Teórico Referencial")
    doc.add_paragraph("Antecedentes de la Investigación")
    doc.add_paragraph(
        "En el plano internacional, García y Pérez (2025) examinaron la eficacia de los sistemas generativos "
        "aplicados a la auditoría de textos académicos, concluyendo que la asistencia algorítmica eleva la "
        "consistencia formal sin alterar la voz original del autor."
    )
    
    # Párrafo con CITA HUÉRFANA INTENCIONAL #1: Martínez (2023)
    doc.add_paragraph(
        "Asimismo, Martínez (2023) destacó la importancia de la experiencia de usuario y la interactividad en "
        "las interfaces de edición asistida, señalando que la previsualización fiel del documento fomenta la "
        "adopción de nuevas plataformas tecnológicas entre investigadores de diversas disciplinas."
    )
    
    doc.add_paragraph("Estructura Jerárquica de Títulos APA")
    doc.add_paragraph("2.1 Procesamiento del Lenguaje Natural en Documentos")
    doc.add_paragraph(
        "El procesamiento del lenguaje natural permite segmentar los documentos en unidades semánticas discretas, "
        "facilitando la identificación de encabezados, párrafos de cuerpo, citas en bloque y tablas informativas."
    )
    
    # TITULO NIVEL 4 (Encabezado con punto y sangría)
    doc.add_paragraph("2.1.1 Arquitectura Basada en Microservicios y FastAPI.")
    doc.add_paragraph(
        "La utilización de microservicios asíncronos permite canalizar las peticiones de formateo sin demoras "
        "perceptibles. Al respecto, Cote (2026) resalta que la arquitectura basada en eventos desacoplados asegura "
        "una alta disponibilidad durante períodos de alta concurrencia académica en plataformas universitarias."
    )
    
    # TITULO NIVEL 5 (Encabezado nivel 5 con punto y sangría)
    doc.add_paragraph("2.1.1.1 Técnicas de Inferencia Semántica con DeepSeek.")
    doc.add_paragraph(
        "La inferencia semántica optimizada analiza el contexto contextual de cada fragmento textual. Como explica "
        "Wilson (2024), la capacidad de clasificar categorías normativas mediante indicaciones estructuradas en JSON "
        "reduce los falsos positivos a niveles insignificantes en comparación con métodos basados puramente en expresiones regulares."
    )
    
    # Párrafo con CITA HUÉRFANA INTENCIONAL #2: López & Gómez (2021)
    doc.add_paragraph(
        "En esta misma línea de análisis, otros estudios de campo sugieren que la combinación de heurísticas y modelos "
        "neuronales minimiza el consumo de recursos de cómputo en servidores distribuidos (López & Gómez, 2021)."
    )
    
    # ─── 5. CAPÍTULO III: MARCO METODOLÓGICO (NIVEL 1 Y 2) ─────────
    doc.add_paragraph("Capítulo III: Marco Metodológico")
    doc.add_paragraph("Diseño de la Investigación")
    doc.add_paragraph(
        "La investigación siguió un enfoque cuantitativo, de nivel explicativo y diseño cuasi-experimental. "
        "Se analizaron 150 documentos técnicos y académicos en formato digital para comparar el desempeño del "
        "sistema automatizado frente al formateo manual tradicional."
    )
    
    doc.add_paragraph("Técnicas de Recolección y Seguridad")
    doc.add_paragraph(
        "Para garantizar la integridad y confidencialidad de la información de los usuarios, se incorporaron "
        "protocolos de escaneo en memoria y sanitización de archivos. Según Rojas (2024), la protección de los "
        "canales de procesamiento contra inyecciones y cargas maliciosas es un requisito indispensable para los "
        "servicios cognitivos modernos desplegados en la nube."
    )
    
    # ─── 6. CAPÍTULO IV: RESULTADOS Y CONCLUSIONES (NIVEL 1 Y 2) ───
    doc.add_paragraph("Capítulo IV: Resultados y Conclusiones")
    doc.add_paragraph("Análisis de Rendimiento")
    doc.add_paragraph(
        "Los ensayos de rendimiento confirmaron que el sistema procesa documentos extensos en un tiempo inferior "
        "a los tres segundos, aplicando sangrías, márgenes de una pulgada, fuentes estandarizadas y alineaciones "
        "justificadas de forma homogénea en la totalidad del manuscrito."
    )
    
    doc.add_paragraph("Conclusiones")
    doc.add_paragraph(
        "Se concluye que la integración sinérgica entre la inteligencia artificial y las reglas deterministas "
        "de estilo APA proporciona una solución robusta y escalable que optimiza el proceso editorial universitario, "
        "garantizando la máxima calidad formal en la presentación de tesis y artículos científicos."
    )
    
    # ─── 7. REFERENCIAS BIBLIOGRÁFICAS (NIVEL 1) ───────────────────
    # Colocadas intencionalmente DESORDENADAS alfabéticamente para comprobar
    # que DocAI las ordena automáticamente de la A a la Z con Sangría Francesa.
    # Contiene:
    # - Citas válidas citadas en el texto: Wilson, Hernández, García, Cote, Navarro, Rojas.
    # - Referencias fantasma NO citadas en el texto: Zúñiga & Morales (2025) y Vargas (2020).
    doc.add_paragraph("Referencias Bibliográficas")
    
    doc.add_paragraph(
        "Zúñiga, R., & Morales, P. (2025). Arquitecturas asíncronas en Python con FastAPI y motores de inteligencia artificial. Editorial Alfaomega."
    )
    doc.add_paragraph(
        "Wilson, K. (2024). The future of academic formatting and large language models in higher education. Journal of Educational Technology, 41(3), 112-129. https://doi.org/10.1016/j.jet.2024.03.008"
    )
    doc.add_paragraph(
        "Hernández, R., Fernández, C., & Baptista, P. (2022). Metodología de la investigación (7a ed.). McGraw-Hill Interamericana."
    )
    doc.add_paragraph(
        "García, J., & Pérez, L. (2025). Inteligencia artificial aplicada a la auditoría documental en la empresa moderna. McGraw-Hill."
    )
    doc.add_paragraph(
        "Cote, M. (2026). Guía avanzada de integración de Odoo y Python para arquitecturas empresariales de alto tráfico. Editorial Tech."
    )
    doc.add_paragraph(
        "Navarro, E. (2021). Manual práctico de normas APA séptima edición para investigadores y tesistas. Ediciones Académicas."
    )
    doc.add_paragraph(
        "Rojas, F. (2024). Seguridad en modelos generativos: prevención de ataques de prompt injection en entornos productivos. Revista Iberoamericana de Ciberseguridad, 12(2), 45-67."
    )
    doc.add_paragraph(
        "Vargas, M. (2020). Automatización de procesos contables en organizaciones sin fines de lucro. Fondo Editorial Universitario."
    )

    # Rutas de destino
    rutas = [
        os.path.join(os.path.dirname(__file__), "test_extenso_docai.docx"),
        os.path.join(os.path.dirname(os.path.dirname(__file__)), "test_extenso_docai.docx"),
        os.path.join(os.path.dirname(os.path.dirname(__file__)), "Documento_Prueba_Normas_APA.docx"),
    ]

    for ruta in rutas:
        doc.save(ruta)
        print(f"[OK] Documento generado exitosamente en: {ruta}")

if __name__ == "__main__":
    generar_documento_prueba()
